from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

C = lambda h: RGBColor.from_string(h)
NAVY, TEAL, ORG, LIGHT, GREY, WHITE = "1B2A41", "0F8B8D", "E07A5F", "EEF4F4", "5B6577", "FFFFFF"
p = Presentation()
p.slide_width = Inches(10)
p.slide_height = Inches(5.625)
blank = p.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        b.fill.solid()
        b.fill.fore_color.rgb = C(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = C(line)
        b.line.width = Pt(1)
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    return b


def text(s, x, y, w, h, runs, size=12, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = t.text_frame
    f.word_wrap = True
    f.vertical_anchor = anchor
    f.margin_left = f.margin_right = Inches(.05)
    f.margin_top = f.margin_bottom = Inches(.03)
    if isinstance(runs, str):
        runs = [runs]
    for i, r in enumerate(runs):
        para = f.paragraphs[0] if i == 0 else f.add_paragraph()
        para.alignment = align
        para.space_after = Pt(3)
        segs = r if isinstance(r, list) else [(r, bold)]
        for tx, b in segs:
            run = para.add_run()
            run.text = tx
            run.font.size = Pt(size)
            run.font.bold = b
            run.font.color.rgb = C(color)
            run.font.name = "Calibri"
    return t


def header(s, n, title):
    box(s, 0, 0, 10, .9, NAVY)
    text(s, .4, .12, 9, .65, title, 24, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.2, 5.3, .6, .3, str(n), 10, GREY, align=PP_ALIGN.RIGHT)


def impact(s, msg):
    box(s, .4, 4.95, 9.2, .38, LIGHT)
    box(s, .4, 4.95, .07, .38, TEAL)
    text(s, .55, 4.95, 9, .38, [[("Why it matters: ", True), (msg, False)]], 11, NAVY, anchor=MSO_ANCHOR.MIDDLE)


def card(s, x, y, w, h, title, body, accent=TEAL, size=11):
    box(s, x, y, w, h, WHITE, "D5DCE4")
    box(s, x, y, w, .06, accent)
    text(s, x + .1, y + .1, w - .2, .35, title, 13, accent, True)
    text(s, x + .1, y + .45, w - .2, h - .5, body, size, NAVY)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def style_table(tb, size, hl_col=None):
    for i, row in enumerate(tb.rows):
        for j, c in enumerate(row.cells):
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = C(NAVY)
            elif hl_col is not None and j == hl_col:
                c.fill.fore_color.rgb = C(LIGHT)
            else:
                c.fill.fore_color.rgb = C(WHITE if i % 2 else "F7FAFA")
            for para in c.text_frame.paragraphs:
                if hl_col is not None and j > 0:
                    para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    run.font.size = Pt(size)
                    run.font.name = "Calibri"
                    run.font.bold = (i == 0 or j == 0 or j == hl_col)
                    run.font.color.rgb = C(WHITE if i == 0 else NAVY)


# ---------- 1 Title
s = p.slides.add_slide(blank)
box(s, 0, 0, 10, 5.625, NAVY)
box(s, 0, 3.55, 10, .05, TEAL)
text(s, .7, .6, 8.6, .4, "22AIE401 · REINFORCEMENT LEARNING · MID REVIEW", 12, "9FD8D9", True)
text(s, .7, 1.1, 8.6, 1.5, "Inverse Reinforcement Learning for Route Preference Learning", 34, WHITE, True)
text(s, .7, 2.7, 8.6, .6, "Learning what taxi drivers value from real GPS trips, then planning routes that match them", 15, "C9D3E0")
text(s, .7, 3.8, 5, 1.2, [[("Abhinav Nair", True), ("   AM.SC.U4AIE23005", False)],
                          [("Advaith K", True), ("   AM.SC.U4AIE23010", False)],
                          [("Akash A", True), ("   AM.SC.U4AIE23011", False)]], 13, WHITE)
text(s, 5.6, 3.8, 3.8, 1.2, [[("RL problem area", True)],
                             "Inverse RL and sequential decision-making: recover the reward behind observed route choices"], 12, "C9D3E0")
text(s, .7, 5.0, 8.6, .4, "Amrita Vishwa Vidyapeetham, Amritapuri", 11, "8FA0B5")
notes(s, "Introduce the team and the problem area: inverse RL. We are not training an agent on a known reward; we recover the reward drivers are implicitly maximising when they pick routes.")

# ---------- 2 Background and statement
s = p.slides.add_slide(blank)
header(s, 2, "Problem Background and Statement")
text(s, .4, 1.05, 4.5, .3, "The real-world problem", 15, TEAL, True)
text(s, .4, 1.4, 4.5, 3.3, [
    "• Drivers rarely take the pure shortest or fastest route. They weigh turns, road type, familiarity and comfort.",
    "• Navigation apps use one fixed cost (distance or time), so routes often feel unnatural.",
    "• A route is a chain of decisions: every junction choice changes the next choices. That is sequential decision-making, an MDP.",
    "• Supervised ML can predict a path but cannot say why drivers chose it."], 12)
box(s, 5.2, 1.05, 4.4, 3.7, LIGHT)
text(s, 5.35, 1.12, 4.1, .3, "Problem statement", 15, TEAL, True)
text(s, 5.35, 1.5, 4.1, 1.6, "Given historical taxi trajectories on a road network, learn the latent reward function that explains driver route choices, and use it to generate routes that match real driver behaviour.", 13, NAVY, True)
text(s, 5.35, 3.2, 4.1, .3, "Why RL?", 13, ORG, True)
text(s, 5.35, 3.5, 4.1, 1.2, "The reward is unknown but the behaviour is observed, so we use inverse RL: actions to reward. The learned weights are readable (what drivers avoid or prefer) and reusable for new origin-destination pairs.", 11.5)
impact(s, "IRL gives an explainable preference model, not just a path predictor.")
notes(s, "Key message: routes are sequential decisions and the reward is hidden. IRL is the natural fit. Our output is interpretable weights.")

# ---------- 3 Existing work
s = p.slides.add_slide(blank)
header(s, 3, "Existing Work, Gap and Proposed Difference")
rows = [("Paper / product", "What it does", "Limitation that remains"),
        ("Ziebart et al. 2008\nMaxEnt IRL", "Learns driver route preferences from GPS with max-entropy IRL", "Small data, basic road features, no context, no open pipeline"),
        ("Barnes et al. 2024\nGoogle Maps", "Large-scale IRL for route quality in production", "Proprietary data and model, cannot be reproduced"),
        ("Zhao & Liang 2023", "Deep IRL with context-dependent rewards", "Black-box network, hard to interpret"),
        ("Zhang 2024, Liu 2025", "Personalised / meta-IRL for driver routes", "Deep models, need many trips per driver"),
        ("Dijkstra / A* (apps)", "Shortest or fastest path on a fixed cost", "No learned preference, same cost for everyone")]
gt = s.shapes.add_table(len(rows), 3, Inches(.4), Inches(1.05), Inches(9.2), Inches(2.7))
tb = gt.table
for j, wd in enumerate([2.3, 3.4, 3.5]):
    tb.columns[j].width = Inches(wd)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        tb.cell(i, j).text = v
style_table(tb, 10.5)
box(s, .4, 3.95, 9.2, .9, LIGHT)
box(s, .4, 3.95, .07, .9, ORG)
text(s, .55, 3.97, 9, .9, [[("Our difference: ", True), ("an exact, fully interpretable linear-reward MaxEnt IRL on open Porto taxi data, with a turn-aware state, a reproducible map-matching pipeline and a live route-comparison dashboard. Next we add familiarity, context (peak, weekend, night) and value-of-time analysis.", False)]], 12)
impact(s, "We do not claim a new IRL algorithm. Our contribution is an open, readable, reproducible study.")
notes(s, "Be honest here: MaxEnt IRL itself is not new. Context-dependent and personalised IRL exist but are deep and black-box. Our angle is interpretability plus open data plus reproducibility.")

# ---------- 4 MDP and algorithm
s = p.slides.add_slide(blank)
header(s, 4, "MDP Design and RL Algorithm")
items = [("Agent", "A taxi driver choosing a route"),
         ("Environment", "Central Porto road graph from OpenStreetMap (~6.5k road segments)"),
         ("State", "(current junction, road just used). The incoming road makes turns Markovian"),
         ("Action", "Choose the next outgoing road segment"),
         ("Reward", "R(s,a) = w . phi(s,a) with 10 features: length, speed, time, road type, turn type, lanes. w is learned"),
         ("Episode / terminal", "One trip, origin to destination. The destination is absorbing and ends the episode")]
for i, (k, v) in enumerate(items):
    y = 1.05 + i * .6
    box(s, .4, y, 1.4, .52, TEAL)
    text(s, .4, y, 1.4, .52, k, 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    box(s, 1.8, y, 3.4, .52, LIGHT)
    text(s, 1.85, y, 3.3, .52, v, 9.5, NAVY, anchor=MSO_ANCHOR.MIDDLE)
card(s, 5.4, 1.05, 4.2, 1.75, "Algorithm: Maximum Entropy IRL",
     ["Soft value iteration gives the policy pi(a|s) ~ exp(Q). Gradient = expert feature counts minus policy feature counts, solved with L-BFGS (L2 = 1e-3).",
      "Why: only demonstrations exist, MaxEnt copes with imperfect drivers, and the graph is small enough for an exact solve, so weights stay readable."], TEAL, 9.5)
card(s, 5.4, 2.9, 4.2, 1.9, "Data: procured and simulated",
     ["Real: Porto taxi GPS (ECML/PKDD 2015), map-matched to OSM. 56% of trips matched; about 3,000 trips from 433 taxis. 80/20 split by origin-destination pair.",
      "Simulated: synthetic drivers with known weights, to check the trainer recovers them."], ORG, 9.5)
impact(s, "Every RL element is defined, and the reward is learned rather than hand-set.")
notes(s, "Walk down the left: agent, environment, state, action, reward, episode. Then justify MaxEnt IRL. Data: real Porto plus synthetic for validation.")

# ---------- 5 System design
from pptx.oxml.ns import qn
from lxml import etree
def arrow(s, x1, y1, x2, y2, color=GREY, dash=False, head=True, w=1.5):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = C(color); ln.line.width = Pt(w)
    el = ln.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(el, qn("a:prstDash")); d.set("val", "dash")
    if head:
        t = etree.SubElement(el, qn("a:tailEnd")); t.set("type", "triangle")
    return ln
def node(s, x, y, w, h, title, sub, fill, size=9):
    box(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, y, w, h, [[(title, True)]] + ([sub] if sub else []), size, WHITE, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
def band(s, y, h, label, x=.45):
    box(s, .4, y, 9.2, h, "F4F7F8", "D5DCE4")
    text(s, x, y + .01, 5.4, .25, label, 9.5, TEAL, True)

s = p.slides.add_slide(blank)
header(s, 5, "System Design Diagram")
# bands
band(s, 1.0, 1.0, "1  DATA AND ENVIRONMENT")
band(s, 2.05, 1.45, "2  RL CORE: MAXIMUM ENTROPY IRL TRAINING LOOP", 1.45)
band(s, 3.6, 1.3, "3  EVALUATION AND SERVING", 2.3)
# band 1 (y 1.27, h .62)
y1, h1 = 1.27, .62
node(s, .5, y1, 1.6, h1, "OSM road graph", "nodes, edges, features φ (10)", NAVY)
node(s, 2.35, y1, 1.7, h1, "Porto taxi GPS", "ECML/PKDD 2015 trips", NAVY)
node(s, 4.3, y1, 1.7, h1, "Map matching", "GPS to road edges (parallel)", NAVY)
node(s, 6.15, y1, 1.7, h1, "Expert paths", "80/20 split by origin-dest", NAVY)
node(s, 8.05, y1, 1.5, h1, "Synthetic drivers", "known w, for validation", GREY)
arrow(s, 4.05, y1 + h1/2, 4.3, y1 + h1/2); arrow(s, 6.0, y1 + h1/2, 6.15, y1 + h1/2)
arrow(s, 8.05, y1 + h1/2, 7.85, y1 + h1/2, GREY, True)
# band 2 (y 2.33, h .7)
y2, h2 = 2.33, .7
node(s, .5, y2, 1.6, h2, "Reward", "R(s,a) = w · φ(s,a)", TEAL)
node(s, 2.35, y2, 1.7, h2, "Soft value iteration", "backward from goal gives π(a|s)", TEAL)
node(s, 4.3, y2, 1.7, h2, "Policy feature counts", "expected visits under π", TEAL)
node(s, 6.15, y2, 1.7, h2, "Gradient", "expert counts − policy counts", TEAL)
node(s, 8.05, y2, 1.5, h2, "L-BFGS update", "new w (L2 = 1e-3)", TEAL)
for a, b in [(2.1, 2.35), (4.05, 4.3), (6.0, 6.15), (7.85, 8.05)]:
    arrow(s, a, y2 + h2/2, b, y2 + h2/2)
arrow(s, 1.3, y1 + h1, 1.3, y2)          # graph -> reward
arrow(s, 7.0, y1 + h1, 7.0, y2)          # expert paths -> gradient
# loop back
arrow(s, 8.8, y2 + h2, 8.8, 3.25, ORG, head=False, w=2)
arrow(s, 8.8, 3.25, 1.3, 3.25, ORG, head=False, w=2)
arrow(s, 1.3, 3.25, 1.3, y2 + h2, ORG, w=2)
text(s, 3.2, 3.05, 4.4, .2, "repeat until converged (about 35 iterations)", 8.5, ORG, True, PP_ALIGN.CENTER)
# band 3
arrow(s, 2.0, 3.25, 2.0, 3.88, TEAL, w=2)
text(s, 2.05, 3.3, 1.3, .2, "learned w*", 8.5, TEAL, True)
node(s, .5, 3.88, 3.0, .45, "Route generation: IRL policy", None, TEAL)
node(s, .5, 4.4, 3.0, .45, "Baselines: Dijkstra, myopic classifier", None, GREY)
node(s, 4.0, 3.95, 2.3, .8, "Metrics", "top-1/3, log-lik, edge F1, Jaccard, length ratio", NAVY)
node(s, 6.8, 3.95, 2.75, .8, "Flask REST API + dashboard", "pick start/end, compare routes on a map", ORG)
arrow(s, 3.5, 4.1, 4.0, 4.3); arrow(s, 3.5, 4.62, 4.0, 4.4)
arrow(s, 6.3, 4.35, 6.8, 4.35)
impact(s, "Raw GPS to interactive route planner runs end to end; the loop learns the reward from data.")
notes(s, "Trace left to right. Band 1 prepares the environment and expert demonstrations (real Porto, or synthetic with known weights as a check). Band 2 is the learning loop: reward from weights, soft value iteration gives the policy, compare its feature counts with the experts', update weights, repeat. Band 3 uses the learned weights to generate routes, scores them against baselines, and serves them through the Flask API to the dashboard.")

# ---------- 6 Novel components
s = p.slides.add_slide(blank)
header(s, 6, "Novel Components of the Project")
nov = [("1", "Open, reproducible Porto IRL pipeline", "GPS, map matching, MaxEnt IRL and evaluation, all on public data with open code.", "Done"),
       ("2", "Turn-aware state, interpretable reward", "State includes the incoming road. Linear 10-feature reward, no black-box network.", "Done"),
       ("3", "Validated trainer", "Synthetic known-reward test recovers weights with cosine similarity 0.987.", "Done"),
       ("4", "Live route-comparison dashboard + API", "Compare IRL vs Dijkstra routes on a map for any origin-destination pair.", "Done"),
       ("5", "Familiarity feature", "Reward for roads drivers know well, computed from training trips only.", "Planned"),
       ("6", "Context-dependent preferences", "Weights shift for peak, weekend and night, with bootstrap confidence intervals.", "Planned"),
       ("7", "Value-of-time analysis", "Metres of detour a driver accepts to save one minute.", "Planned")]
for i, (n, t, b, st) in enumerate(nov):
    y = 1.02 + i * .55
    col = TEAL if st == "Done" else ORG
    box(s, .4, y, .45, .48, col)
    text(s, .4, y, .45, .48, n, 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    box(s, .85, y, 8.75, .48, LIGHT)
    text(s, .92, y, 3.4, .48, t, 10.5, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.35, y, 4.4, .48, b, 9.5, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 8.8, y + .1, .7, .28, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 8.8, y + .1, .7, .28, st, 9, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
impact(s, "Items 5 to 7 are the core extension that goes beyond plain MaxEnt IRL.")
notes(s, "Be clear which are done and which are planned. Items 5-7 are the research contribution still to complete.")

# ---------- 7 Demo and next plan
s = p.slides.add_slide(blank)
header(s, 7, "Demo So Far and Next Plan")
s.shapes.add_picture("midreview_assets/curve.png", Inches(.4), Inches(1.0), width=Inches(3.0))
s.shapes.add_picture("midreview_assets/weights.png", Inches(.4), Inches(2.95), width=Inches(3.0))
rows = [("Held-out Porto trips", "IRL", "Classifier", "Shortest", "Fastest"),
        ("Top-1 accuracy", "92.4%", "81.1%", "-", "-"),
        ("Top-3 accuracy", "99.9%", "99.7%", "-", "-"),
        ("Log-likelihood", "-0.18", "-0.44", "-", "-"),
        ("Edge F1", "0.61", "0.09", "0.64", "0.45"),
        ("Path length ratio", "0.92", "2.10", "0.88", "1.01")]
gt = s.shapes.add_table(6, 5, Inches(3.6), Inches(1.0), Inches(6.0), Inches(1.9))
tb = gt.table
for j, wd in enumerate([1.8, .9, 1.1, 1.1, 1.1]):
    tb.columns[j].width = Inches(wd)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        tb.cell(i, j).text = v
style_table(tb, 10, hl_col=1)
text(s, 3.6, 3.0, 6, .3, "What works now", 12, TEAL, True)
text(s, 3.6, 3.28, 6, 1.0, ["• Pipeline, trainer, baselines, Flask API and map dashboard run end to end.",
                            "• IRL beats the ML classifier on next-road prediction and is far closer to real routes than fastest-path.",
                            "• Weights: drivers avoid sharp turns (-1.48) and long travel time (-1.26)."], 10.5)
text(s, 3.6, 4.2, 6, .3, "Honest note and remaining work", 12, ORG, True)
text(s, 3.6, 4.45, 6, .5, "Plain shortest path still scores slightly higher edge F1 (0.64 vs 0.61). Next: familiarity, context, value of time, 3 seeds, taxi-level split, popularity-weighted Dijkstra, final report.", 10)
box(s, .4, 4.95, 9.2, .38, LIGHT)
box(s, .4, 4.95, .07, .38, TEAL)
text(s, .55, 4.95, 9, .38, [[("Why it matters: ", True), ("about 50% done. Pipeline, results and demo work; the novel extensions remain.", False)]], 11, NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Run the dashboard live: python run_api.py, then open dashboard/index.html and pick a start and end point. Show the training curve (NLL 10.6 to 6.40), then the table. Mention the shortest-path caveat openly.")

p.save("RL_MidReview_Final.pptx")
