from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

C = lambda h: RGBColor.from_string(h)
INK, MUTED, ACC, SLATE, OCHRE = "262626", "6B6B6B", "8A2432", "3E5C76", "A8742A"
PAPER, BAND, LINE, WHITE = "FAF8F4", "F1EDE5", "CFC8BB", "FFFFFF"
SERIF, SANS = "Georgia", "Calibri"
p = Presentation()
p.slide_width = Inches(10)
p.slide_height = Inches(5.625)
blank = p.slide_layouts[6]


def new_slide():
    s = p.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C(PAPER)
    return s


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=.75, dash=False):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        b.fill.solid()
        b.fill.fore_color.rgb = C(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = C(line)
        b.line.width = Pt(lw)
        if dash:
            d = etree.SubElement(b.line._get_or_add_ln(), qn("a:prstDash"))
            d.set("val", "dash")
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    return b


def text(s, x, y, w, h, runs, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=SANS, italic=False, after=3):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = t.text_frame
    f.word_wrap = True
    f.vertical_anchor = anchor
    f.margin_left = f.margin_right = Inches(.05)
    f.margin_top = f.margin_bottom = Inches(.03)
    if isinstance(runs, str):
        runs = [runs]
    for i, r in enumerate(runs):
        para = f.paragraphs[0] if i == 0 else f.add_paragraph()
        para.alignment = align
        para.space_after = Pt(after)
        segs = r if isinstance(r, list) else [(r, bold)]
        for tx, b in segs:
            run = para.add_run()
            run.text = tx
            run.font.size = Pt(size)
            run.font.bold = b
            run.font.italic = italic
            run.font.color.rgb = C(color)
            run.font.name = font
    return t


def hline(s, x, y, w, color=LINE, weight=.75):
    l = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    l.line.color.rgb = C(color)
    l.line.width = Pt(weight)


def header(s, n, title):
    text(s, .5, .28, 9, .55, title, 26, INK, False, anchor=MSO_ANCHOR.MIDDLE, font=SERIF)
    hline(s, .5, .9, 9, LINE)
    box(s, .5, .88, .9, .04, ACC)
    text(s, 9.0, 5.28, .5, .25, str(n), 9, MUTED, align=PP_ALIGN.RIGHT)


def takeaway(s, msg):
    hline(s, .5, 4.95, 9, LINE)
    text(s, .5, 4.98, 8.4, .32, msg, 11, MUTED, italic=True, anchor=MSO_ANCHOR.MIDDLE, font=SERIF)


def card(s, x, y, w, h, title, body, accent=ACC, size=10.5):
    box(s, x, y, w, h, WHITE, LINE)
    box(s, x, y, .05, h, accent)
    text(s, x + .15, y + .08, w - .25, .3, title, 12.5, accent, True, font=SERIF)
    text(s, x + .15, y + .42, w - .25, h - .48, body, size, INK)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def style_table(tb, size, hl_col=None):
    for i, row in enumerate(tb.rows):
        for j, c in enumerate(row.cells):
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = C("E6E0D3")
            elif hl_col is not None and j == hl_col:
                c.fill.fore_color.rgb = C(BAND)
            else:
                c.fill.fore_color.rgb = C(WHITE)
            for para in c.text_frame.paragraphs:
                if hl_col is not None and j > 0:
                    para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    run.font.size = Pt(size)
                    run.font.name = SANS
                    run.font.bold = (i == 0 or j == 0 or j == hl_col)
                    run.font.color.rgb = C(INK)


def arrow(s, x1, y1, x2, y2, color=MUTED, dash=False, head=True, w=1.25, both=False):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = C(color)
    ln.line.width = Pt(w)
    el = ln.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(el, qn("a:prstDash"))
        d.set("val", "dash")
    if both:
        h = etree.SubElement(el, qn("a:headEnd"))
        h.set("type", "triangle")
    if head:
        t = etree.SubElement(el, qn("a:tailEnd"))
        t.set("type", "triangle")
    return ln


# ---------- 1 Title
s = new_slide()
box(s, 0, 0, .22, 5.625, ACC)
text(s, .9, .7, 8.4, .35, "22AIE401  ·  Reinforcement Learning  ·  Mid Review", 12, ACC, True)
text(s, .9, 1.2, 8.4, 1.6, "Inverse Reinforcement Learning for Route Preference Learning", 36, INK, False, font=SERIF)
hline(s, .9, 2.95, 1.2, ACC, 2.5)
text(s, .9, 3.1, 8.2, .6, "Learning what taxi drivers value from real GPS trips, then planning routes that match them.", 15, MUTED, font=SERIF, italic=True)
text(s, .9, 3.95, 4.3, 1.2, [[("Abhinav Nair", True), ("   AM.SC.U4AIE23005", False)],
                              [("Advaith K", True), ("   AM.SC.U4AIE23010", False)],
                              [("Akash A", True), ("   AM.SC.U4AIE23011", False)]], 13, INK)
text(s, 5.5, 3.95, 3.8, 1.2, [[("Problem area", True)],
                              "Inverse RL and sequential decision-making: recover the reward behind observed route choices."], 12, MUTED)
text(s, .9, 5.15, 8, .3, "Amrita Vishwa Vidyapeetham, Amritapuri", 10, MUTED)
notes(s, "Introduce the team and the problem area: inverse RL. We are not training an agent on a known reward; we recover the reward drivers are implicitly maximising when they pick routes.")

# ---------- 2 Background and statement
s = new_slide()
header(s, 2, "Problem background and statement")
text(s, .5, 1.1, 4.4, .3, "The real-world problem", 15, ACC, False, font=SERIF)
text(s, .5, 1.5, 4.4, 3.3, [
    "• Drivers rarely take the pure shortest or fastest route. They weigh turns, road type, familiarity and comfort.",
    "• Navigation apps use one fixed cost (distance or time), so routes often feel unnatural.",
    "• A route is a chain of decisions: every junction choice changes the next choices. That is sequential decision-making, an MDP.",
    "• Supervised ML can predict a path but cannot say why drivers chose it."], 12, after=6)
box(s, 5.2, 1.1, 4.3, 3.65, BAND)
box(s, 5.2, 1.1, .05, 3.65, ACC)
text(s, 5.4, 1.18, 4.0, .3, "Problem statement", 15, ACC, False, font=SERIF)
text(s, 5.4, 1.58, 4.0, 1.6, "Given historical taxi trajectories on a road network, learn the latent reward function that explains driver route choices, and use it to generate routes that match real driver behaviour.", 13, INK, font=SERIF)
text(s, 5.4, 3.3, 4.0, .3, "Why RL?", 13, ACC, True)
text(s, 5.4, 3.6, 4.0, 1.1, "The reward is unknown but the behaviour is observed, so we use inverse RL: from actions back to reward. The learned weights are readable and reusable for new origin-destination pairs.", 11.5)
takeaway(s, "IRL gives an explainable preference model, not just a path predictor.")
notes(s, "Key message: routes are sequential decisions and the reward is hidden. IRL is the natural fit. Our output is interpretable weights.")

# ---------- 3 Existing work
s = new_slide()
header(s, 3, "Existing work, gap and proposed difference")
rows = [("Paper / product", "What it does", "Limitation that remains"),
        ("Ziebart et al. 2008, MaxEnt IRL", "Learns driver route preferences from GPS with max-entropy IRL", "Small data, basic road features, no context, no open pipeline"),
        ("Barnes et al. 2024, Google Maps", "Large-scale IRL for route quality in production", "Proprietary data and model, cannot be reproduced"),
        ("Zhao & Liang 2023", "Deep IRL with context-dependent rewards", "Black-box network, hard to interpret"),
        ("Zhang 2024, Liu 2025", "Personalised / meta-IRL for driver routes", "Deep models, need many trips per driver"),
        ("Dijkstra / A* (navigation apps)", "Shortest or fastest path on a fixed cost", "No learned preference, same cost for everyone")]
gt = s.shapes.add_table(len(rows), 3, Inches(.5), Inches(1.1), Inches(9), Inches(2.6))
tb = gt.table
for j, wd in enumerate([2.6, 3.3, 3.1]):
    tb.columns[j].width = Inches(wd)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        tb.cell(i, j).text = v
style_table(tb, 10.5)
box(s, .5, 3.95, 9, .9, BAND)
box(s, .5, 3.95, .05, .9, ACC)
text(s, .68, 3.98, 8.7, .85, [[("Our difference: ", True), ("an exact, fully interpretable linear-reward MaxEnt IRL on open Porto taxi data, with a turn-aware state, a reproducible map-matching pipeline and a live route-comparison dashboard. Next we add familiarity, context (peak, weekend, night) and value-of-time analysis.", False)]], 12)
takeaway(s, "We do not claim a new IRL algorithm. Our contribution is an open, readable, reproducible study.")
notes(s, "Be honest here: MaxEnt IRL itself is not new. Context-dependent and personalised IRL exist but are deep and black-box. Our angle is interpretability plus open data plus reproducibility.")

# ---------- 4 MDP and algorithm
s = new_slide()
header(s, 4, "MDP design and RL algorithm")
items = [("Agent", "A taxi driver choosing a route"),
         ("Environment", "Central Porto road graph from OpenStreetMap (~6.5k road segments)"),
         ("State", "(current junction, road just used). The incoming road makes turns Markovian"),
         ("Action", "Choose the next outgoing road segment"),
         ("Reward", "R(s,a) = w · φ(s,a) with 10 features: length, speed, time, road type, turn type, lanes. w is learned"),
         ("Episode / terminal", "One trip, origin to destination. The destination is absorbing and ends the episode")]
for i, (k, v) in enumerate(items):
    y = 1.1 + i * .6
    hline(s, .5, y + .55, 4.8, LINE)
    text(s, .5, y, 1.35, .52, k, 12, ACC, False, anchor=MSO_ANCHOR.MIDDLE, font=SERIF)
    text(s, 1.85, y, 3.45, .52, v, 9.5, INK, anchor=MSO_ANCHOR.MIDDLE)
card(s, 5.6, 1.1, 3.9, 2.15, "Algorithm: Maximum Entropy IRL",
     ["Soft value iteration (80-step horizon, no discounting) gives π(a|s) ∝ exp(Q). Gradient = expert minus policy feature counts, solved with L-BFGS (L2 = 1e-3).",
      "Why IRL, not Q-learning or PPO: no reward is given, only driver choices, so we learn the reward first.",
      "Transitions are deterministic: choosing a road leads to its next junction."], ACC, 9.5)
card(s, 5.6, 3.35, 3.9, 1.5, "Data: procured and simulated",
     ["Real: Porto taxi GPS, map-matched to OSM. 56% matched, about 3,000 trips from 433 taxis, split 80/20 by origin-destination pair.",
      "Simulated: synthetic drivers with known weights, to check the trainer."], SLATE, 9.5)
takeaway(s, "Every RL element is defined, and the reward is learned rather than hand-set.")
notes(s, "Walk down the left: agent, environment, state, action, reward, episode. Algorithm: this is inverse RL, so there is no reward to optimise; we learn it. Soft value iteration is the forward-RL step inside each update: it runs backward from the destination over an 80-step horizon with no discounting (gamma = 1) and gives the stochastic policy. Transitions are deterministic because a road segment leads to one next junction, so all randomness is in the policy. Why not Q-learning or PPO: they need a reward function, and we only have demonstrations. Data: real Porto plus synthetic for validation.")


# ---------- 4b MDP illustration
s = new_slide()
header(s, 5, "Route choice as an MDP")
w0, h0 = Image.open("midreview_assets/mdp_crop.png").size
H = 3.95
W = H * w0 / h0
px = 9.5 - W
box(s, px - .06, 1.04, W + .12, H + .12, WHITE, LINE)
s.shapes.add_picture("midreview_assets/mdp_crop.png", Inches(px), Inches(1.1), height=Inches(H))
rows5 = [("State", "(junction, incoming road)  ·  6,644", "One per directed road segment; the incoming road tells a turn from a U-turn"),
         ("Action", "next outgoing road  ·  14,788", "Every road-to-next-road pair; 2.2 choices per state on average, at most 5"),
         ("Transition", "deterministic, P = 1.0", "A road leads to exactly one next junction; the randomness is in the policy"),
         ("Reward", "R(s,a) = w · φ(s,a), 10 features", "Learned. Mean R: straight −0.58, slight turn −1.14, sharp turn −2.24"),
         ("Terminal", "destination, absorbing", "One trip is one episode: 5 to 75 roads, mean 33 in our data"),
         ("Horizon", "H = 80, no discounting", "Above the 75-road cap on trips, and every road counts equally")]
for i, (k, d, why) in enumerate(rows5):
    y = 1.08 + i * .64
    hline(s, .5, y + .62, px - 1.0, LINE)
    text(s, .5, y, 1.05, .6, k, 13, ACC, False, anchor=MSO_ANCHOR.MIDDLE, font=SERIF)
    text(s, 1.55, y + .02, px - 2.1, .26, d, 10.5, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.55, y + .27, px - 2.1, .33, why, 8.5, MUTED)
text(s, px, 1.1 + H + .1, W, .3, "Illustrative trip from the Phase 2 report", 9, MUTED, italic=True)
notes(s, "Figure 2 from the Phase 2 report: an illustrative MDP for one trip, with a wait, a bypass, a shortcut, a refuel branch, a successful arrival (S4) and a blocked-road trap (S5). Circles are states, diamonds are actions, arrows carry transition probability and reward. In our real model: 6,644 states (one per directed road segment of the Porto graph), 14,788 actions (road-to-next-road pairs, 2.2 per state on average), deterministic transitions, 10-feature reward learned by MaxEnt IRL. Mean learned reward by turn type: straight -0.58, slight -1.14, sharp -2.24, and 73 percent of real driver steps are straight. Horizon 80 sits just above the 75-road cap on trips.")

# ---------- 5 System design
def mini(s, x, y, w, h, icon, title, sub, accent=INK, ts=9.5):
    box(s, x, y, w, h, WHITE, accent, lw=1)
    isz = .3
    s.shapes.add_picture("midreview_assets/icons/%s.png" % icon, Inches(x + w / 2 - isz / 2), Inches(y + .04), width=Inches(isz))
    text(s, x, y + .33, w, .22, title, ts, accent, True, PP_ALIGN.CENTER, font=SERIF)
    if sub:
        text(s, x, y + .54, w, .2, sub, 8, MUTED, False, PP_ALIGN.CENTER)


def layer(s, y, h, num, name, sub, accent, tint):
    box(s, .5, y, 9, h, tint)
    box(s, .5, y, 1.5, h, accent)
    text(s, .5, y + h / 2 - .45, 1.5, .5, num, 26, WHITE, False, PP_ALIGN.CENTER, font=SERIF, after=0)
    text(s, .5, y + h / 2 + .02, 1.5, .28, name, 12.5, WHITE, True, PP_ALIGN.CENTER, font=SERIF, after=0)
    text(s, .5, y + h / 2 + .3, 1.5, .22, sub, 9, "EADCDD" if accent == ACC else "DDE6EE" if accent == SLATE else "DADADA", False, PP_ALIGN.CENTER, italic=True, after=0)


s = new_slide()
header(s, 6, "System design")
layer(s, .98, 1.14, "3", "Application", "online", SLATE, "E4EAF0")
layer(s, 2.45, 1.3, "2", "Learning", "offline", ACC, "F3E6E4")
layer(s, 4.08, 1.12, "1", "Data", "offline", INK, "EFEBE2")
GX = [2.3 + i * 1.45 for i in range(5)]
W = 1.15
H = .78
# ---- layer 3: application
ya = 1.17
mini(s, 2.3, ya, 1.9, H, "monitor", "Web dashboard", "pick start and end", SLATE)
mini(s, 4.9, ya, 1.9, H, "code", "Flask REST API", "/api/route", SLATE)
mini(s, 7.9, ya, 1.5, H, "fork", "Route planner", "IRL + baselines", SLATE)
arrow(s, 4.2, ya + .4, 4.9, ya + .4, SLATE, both=True, w=1.5)
arrow(s, 6.8, ya + .4, 7.9, ya + .4, SLATE, both=True, w=1.5)
# ---- layer 2: learning
yl = 2.6
mini(s, GX[0], yl, W, H, "target", "Reward", "w · φ(s,a)", ACC)
mini(s, GX[1], yl, W, H, "layers", "Soft value iter.", "policy π(a|s)", ACC, 9)
mini(s, GX[2], yl, W, H, "gradient", "Gradient", "expert − policy", ACC)
mini(s, GX[3], yl, W, H, "cycle", "L-BFGS", "update w", ACC)
mini(s, GX[4], yl, W, H, "db_acc", "Weights w*", "saved model", ACC)
for i in range(4):
    arrow(s, GX[i] + W, yl + .4, GX[i + 1], yl + .4, MUTED, w=1.5)
lyy = 3.55
arrow(s, GX[3] + W / 2, yl + H, GX[3] + W / 2, lyy, ACC, head=False, w=1.3)
arrow(s, GX[3] + W / 2, lyy, GX[0] + .88, lyy, ACC, head=False, w=1.3)
arrow(s, GX[0] + .88, lyy, GX[0] + .88, yl + H, ACC, w=1.3)
text(s, 3.4, lyy - .01, 3.2, .2, "repeat about 35 times", 8.5, ACC, False, PP_ALIGN.CENTER, italic=True, after=0)
# ---- layer 1: data
yd = 4.25
mini(s, GX[0], yd, W, H, "pin_ink", "Porto GPS", "+ OSM roads", INK)
mini(s, GX[1], yd, W, H, "snap_ink", "Map matching", "GPS to edges", INK, 9)
mini(s, GX[2], yd, W, H, "route_ink", "Expert routes", "80 / 20 split", INK, 9)
mini(s, GX[3], yd, W, H, "dice", "Synthetic", "optional", MUTED)
arrow(s, GX[0] + W, yd + .4, GX[1], yd + .4, MUTED, w=1.5)
arrow(s, GX[1] + W, yd + .4, GX[2], yd + .4, MUTED, w=1.5)
arrow(s, GX[3], yd + .4, GX[2] + W, yd + .4, MUTED, w=1.2)
# ---- arrows between layers (data flows upward)
arrow(s, GX[2] + W / 2, yd, GX[2] + W / 2, yl + H + .0 + .0, MUTED, w=1.6)
text(s, GX[2] + W / 2 + .08, 3.82, 1.4, .2, "expert routes", 8.5, MUTED, False, italic=True, after=0)
arrow(s, GX[0] + .28, yd, GX[0] + .28, yl + H, MUTED, w=1.6)
text(s, GX[0] + .36, 3.82, 1.6, .2, "road features φ", 8.5, MUTED, False, italic=True, after=0)
arrow(s, GX[4] + W / 2, yl, GX[4] + W / 2, ya + H, SLATE, w=1.8)
text(s, GX[4] + W / 2 + .08, 2.17, .9, .2, "w*", 9, SLATE, True, italic=True, after=0)
notes(s, "Read it bottom to top. Layer 1, data: Porto GPS trips are map-matched onto the OSM road graph into expert routes (80/20 split by origin-destination pair); synthetic drivers with a known reward are an optional check. Layer 2, learning: reward w.phi, soft value iteration gives the policy, the gradient compares policy and expert feature counts, L-BFGS updates w, repeat about 35 times; the saved weights w* are tested on held-out trips. Layer 3, application: the dashboard calls the Flask API, which asks the route planner for the IRL route and the Dijkstra and classifier baselines using w* and the road graph.")

# ---------- 6 Novel components
s = new_slide()
header(s, 7, "Novel components of the project")
nov = [("1", "Open, reproducible Porto IRL pipeline", "GPS, map matching, MaxEnt IRL and evaluation, all on public data with open code.", "Done"),
       ("2", "Turn-aware state, interpretable reward", "State includes the incoming road. Linear 10-feature reward, no black-box network.", "Done"),
       ("3", "Validated trainer", "Synthetic known-reward test recovers weights with cosine similarity 0.987.", "Done"),
       ("4", "Live route-comparison dashboard + API", "Compare IRL vs Dijkstra routes on a map for any origin-destination pair.", "Done"),
       ("5", "Familiarity feature", "Reward for roads drivers know well, computed from training trips only.", "Planned"),
       ("6", "Context-dependent preferences", "Weights shift for peak, weekend and night, with bootstrap confidence intervals.", "Planned"),
       ("7", "Value-of-time analysis", "Metres of detour a driver accepts to save one minute.", "Planned")]
for i, (n, t, b, st) in enumerate(nov):
    y = 1.05 + i * .55
    done = st == "Done"
    col = SLATE if done else ACC
    hline(s, .5, y + .52, 9, LINE)
    text(s, .5, y, .45, .5, n, 18, col, False, anchor=MSO_ANCHOR.MIDDLE, font=SERIF)
    text(s, .95, y, 3.4, .5, t, 11, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.4, y, 4.2, .5, b, 9.5, INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.6, y, .9, .5, st, 10.5, col, True, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, italic=not done)
takeaway(s, "Items 5 to 7 are the core extension that goes beyond plain MaxEnt IRL.")
notes(s, "Be clear which are done and which are planned. Items 5-7 are the research contribution still to complete.")

# ---------- 7 Demo and next plan
s = new_slide()
header(s, 8, "Demo so far and next plan")
s.shapes.add_picture("midreview_assets/curve.png", Inches(.5), Inches(1.05), width=Inches(3.0))
s.shapes.add_picture("midreview_assets/weights.png", Inches(.5), Inches(2.95), width=Inches(3.0))
rows = [("Held-out Porto trips", "IRL", "Classifier", "Shortest", "Fastest"),
        ("Top-1 accuracy", "92.4%", "81.1%", "-", "-"),
        ("Top-3 accuracy", "99.9%", "99.7%", "-", "-"),
        ("Log-likelihood", "-0.18", "-0.44", "-", "-"),
        ("Edge F1", "0.61", "0.09", "0.64", "0.45"),
        ("Path length ratio", "0.92", "2.10", "0.88", "1.01")]
gt = s.shapes.add_table(6, 5, Inches(3.7), Inches(1.05), Inches(5.8), Inches(1.85))
tb = gt.table
for j, wd in enumerate([1.7, .9, 1.1, 1.05, 1.05]):
    tb.columns[j].width = Inches(wd)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        tb.cell(i, j).text = v
style_table(tb, 10, hl_col=1)
text(s, 3.7, 3.0, 5.8, .3, "What works now", 13, SLATE, False, font=SERIF)
text(s, 3.7, 3.3, 5.8, 1.0, ["• Pipeline, trainer, baselines, Flask API and map dashboard run end to end.",
                             "• IRL beats the ML classifier on next-road prediction and is far closer to real routes than fastest-path.",
                             "• Weights: drivers avoid sharp turns (-1.48) and long travel time (-1.26)."], 10.5)
text(s, 3.7, 4.2, 5.8, .3, "Honest note and remaining work", 13, ACC, False, font=SERIF)
text(s, 3.7, 4.47, 5.8, .5, "Plain shortest path still scores slightly higher edge F1 (0.64 vs 0.61). Next: familiarity, context, value of time, 3 seeds, taxi-level split, popularity-weighted Dijkstra, final report.", 10)
takeaway(s, "About 50% done. Pipeline, results and demo work; the novel extensions remain.")
notes(s, "Run the dashboard live: python run_api.py, then open dashboard/index.html and pick a start and end point. Show the training curve (NLL 10.6 to 6.40), then the table. Mention the shortest-path caveat openly.")

p.save("RL_MidReview_Final.pptx")
