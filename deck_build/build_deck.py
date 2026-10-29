from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION

NAVY = RGBColor(0x0B, 0x1F, 0x3A); TEAL = RGBColor(0x0F, 0x76, 0x6E); AMBER = RGBColor(0xF5, 0x9E, 0x0B)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB); WHITE = RGBColor(255, 255, 255); TXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x64, 0x74, 0x8B); GREEN = RGBColor(0x15, 0x80, 0x3D); RED = RGBColor(0xB9, 0x1C, 0x1C)
BLUE = RGBColor(0x25, 0x63, 0xEB); CARD = RGBColor(0xFF, 0xFF, 0xFF); LINE = RGBColor(0xD6, 0xDE, 0xEA)
HF, BF = 'Cambria', 'Calibri'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
N = [0]


def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def tb(slide, x, y, w, h, text, size=16, bold=False, color=TXT, font=BF, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = s.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = l; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = font; r.font.italic = italic
    return s


def bullets(slide, x, y, w, h, items, size=16, color=TXT, gap=6):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = s.text_frame
    tf.word_wrap = True; tf.margin_left = Inches(0.05)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, sep, rest = it.partition('::')
        rest = rest.rstrip()
        if sep and not rest:
            it = head
            sep = ''
        if rest:
            r = p.add_run(); r.text = '▪ ' + head; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = TEAL; r.font.name = BF
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = BF
        else:
            r = p.add_run(); r.text = '▪ ' + it; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = BF
    return s


def card(slide, x, y, w, h, fill=CARD, line=LINE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.06; s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(1); s.shadow.inherit = False
    return s


def new(title, section=None, notes=''):
    s = prs.slides.add_slide(BLANK); bg(s, LIGHT); N[0] += 1
    if section:
        tb(s, 0.6, 0.35, 8, 0.3, section.upper(), 12, True, TEAL)
    tb(s, 0.6, 0.62, 12.1, 0.8, title, 28, True, NAVY, HF)
    tb(s, 12.2, 7.05, 0.8, 0.3, str(N[0]), 11, False, MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def stat(slide, x, y, w, big, label, sub='', color=TEAL, bigsize=34):
    card(slide, x, y, w, 1.55)
    tb(slide, x, y + 0.12, w, 0.7, big, bigsize, True, color, HF, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(slide, x, y + 0.82, w, 0.35, label, 14, True, TXT, BF, PP_ALIGN.CENTER)
    if sub:
        tb(slide, x, y + 1.14, w, 0.3, sub, 11, False, MUTED, BF, PP_ALIGN.CENTER)


def table(slide, x, y, w, rows, colw, size=12, rowh=0.36, header=NAVY, hl_row=None):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    t = shp.table
    tblPr = shp._element.graphic.graphicData.tbl.tblPr
    tblPr.set('bandRow', '0'); tblPr.set('firstRow', '0')
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = BF
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = header; r.font.bold = True; r.font.color.rgb = WHITE
            else:
                c.fill.fore_color.rgb = RGBColor(0xE6, 0xF4, 0xF1) if i == hl_row else (WHITE if i % 2 else RGBColor(0xF1, 0xF5, 0xF9))
                r.font.color.rgb = TXT; r.font.bold = (i == hl_row)
    return t


def bar_chart(slide, x, y, w, h, cats, vals, title='', horizontal=True, colors=None, fmt='0.000', vmin=None, vmax=None, size=12):
    cd = CategoryChartData(); cd.categories = cats; cd.add_series('v', vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart; ch.has_legend = False
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = title
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.bold = True; r.font.name = BF; r.font.color.rgb = NAVY
    pl = ch.plots[0]; pl.gap_width = 55; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = fmt; dl.number_format_is_linked = False; dl.font.size = Pt(size); dl.font.name = BF
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(size); ca.tick_labels.font.name = BF; ca.format.line.color.rgb = LINE
    ca.has_major_gridlines = False
    if horizontal:
        ca.reverse_order = True
    va = ch.value_axis; va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = LINE
    va.tick_labels.font.size = Pt(10); va.tick_labels.font.name = BF; va.format.line.fill.background()
    va.tick_labels.number_format = fmt; va.tick_labels.number_format_is_linked = False
    if vmin is not None: va.minimum_scale = vmin
    if vmax is not None: va.maximum_scale = vmax
    ser = pl.series[0]; ser.invert_if_negative = False
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    for i in range(len(vals)):
        pt = ser.points[i]; pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = (colors[i] if colors else TEAL)
    from lxml import etree
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    for dpt in ser._element.findall('{%s}dPt' % ns):
        if dpt.find('{%s}invertIfNegative' % ns) is None:
            inv = etree.SubElement(dpt, '{%s}invertIfNegative' % ns); inv.set('val', '0')
            idx = dpt.find('{%s}idx' % ns); dpt.remove(inv); idx.addnext(inv)
    return ch


# ================= 1. TITLE =================
s = prs.slides.add_slide(BLANK); bg(s, NAVY); N[0] += 1
tb(s, 0.8, 1.2, 11.5, 0.4, '22AIE401 · REINFORCEMENT LEARNING · AMRITA VISHWA VIDYAPEETHAM, AMRITAPURI', 13, True, AMBER)
tb(s, 0.8, 1.85, 11.7, 2.0, ['Why Nobody Drives the Shortest Route:', 'Learning Taxi Route Preferences with Maximum Entropy Inverse RL'], 38, True, WHITE, HF)
tb(s, 0.8, 4.1, 11.5, 0.9, 'Recovering interpretable driver preferences from real Porto taxi GPS trips on an OpenStreetMap road graph', 20, False, RGBColor(0xCB, 0xD5, 0xE1))
for i, (b, l) in enumerate([('3,000', 'real map-matched taxi trips'), ('11', 'interpretable reward features'), ('0.682', 'edge F1 vs 0.636 for shortest path'), ('93.7%', 'top-1 junction accuracy')]):
    x = 0.8 + i * 3.05
    tb(s, x, 5.45, 2.9, 0.6, b, 32, True, AMBER, HF)
    tb(s, x, 6.05, 2.8, 0.6, l, 13, False, RGBColor(0xCB, 0xD5, 0xE1))
s.notes_slide.notes_text_frame.text = ("Open with the hook: nobody drives the shortest route, so what are drivers optimising? We learn their hidden reward from real GPS data using Maximum Entropy Inverse RL. "
    "Headline: on 600 held-out Porto trips our model reaches edge F1 0.682 vs 0.636 for Dijkstra shortest path, and predicts the next road with 93.7% top-1 accuracy. We will be honest about where the gain is modest.")

# ================= 2. PROBLEM =================
s = new('Apps assume one metric; real drivers do not', 'Problem & motivation',
        "Motivate. Shortest and fastest routes optimise one fixed number. Real taxi drivers trade off distance, turns, road familiarity. Those preferences are not written anywhere: that is why we need Inverse RL — observe behaviour, recover the reward.")
card(s, 0.6, 1.6, 6.0, 4.9)
tb(s, 0.85, 1.75, 5.5, 0.4, 'The gap', 20, True, NAVY, HF)
bullets(s, 0.85, 2.3, 5.5, 4.0, ['Shortest / fastest path:: one fixed metric, no driver comfort or habits', 'Preferences are hidden:: drivers cannot articulate their own rules',
        'Inverse RL:: observe demonstrations, recover the reward that explains them', 'Why it matters:: personalised navigation, realistic simulation, human-compatible autonomy'], 16)
card(s, 6.9, 1.6, 5.85, 4.9)
tb(s, 7.15, 1.75, 5.4, 0.4, 'Our task', 20, True, NAVY, HF)
tb(s, 7.15, 2.3, 5.4, 1.0, 'Given a road graph G and N map-matched trips D = {τ₁…τₙ}, find reward weights w that make the observed routes most likely:', 15, False, TXT)
card(s, 7.15, 3.5, 5.35, 1.0, fill=RGBColor(0xEC, 0xF3, 0xFB), line=RGBColor(0xBF, 0xD4, 0xEE))
tb(s, 7.25, 3.6, 5.2, 0.8, 'w* = argmax_w Σᵢ log P(τᵢ | w, oᵢ, dᵢ) − λ‖w‖²', 17, True, NAVY, 'Cambria', PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
bullets(s, 7.15, 4.75, 5.4, 1.7, ['Forward RL:: reward → policy', 'Inverse RL (ours):: demonstrations → reward'], 16)

# ================= 3. RELATED WORK =================
s = new('Related work: from feature matching to deep IRL', 'Literature review',
        "Walk through the table quickly. Ziebart AAAI 2008 is our base method (MaxEnt IRL on Pittsburgh taxis, private data). Abbeel & Ng ICML 2004 is feature matching. Wulfmeier 2015 and Finn ICML 2016 are non-linear, sample-based extensions. McFadden logit is the econometric route-choice tradition; Dijkstra/A* is our baseline. Newson & Krumm is our map-matching. Venue details should be double-checked against the papers.")
rows = [['Work', 'Venue', 'Idea', 'Limit for our setting'],
        ['Ziebart et al.', 'AAAI 2008, Chicago', 'MaxEnt IRL on taxi GPS, linear reward', 'Private data, not reproducible'],
        ['Ziebart et al.', 'UbiComp 2008, Seoul', 'Predict route, turn and destination with context', 'Private data; destination inference'],
        ['Abbeel & Ng', 'ICML 2004, Banff', 'Feature-expectation matching (apprenticeship)', 'Mixture of policies; ambiguous'],
        ['Wulfmeier et al.', 'NIPS 2015 Deep RL workshop', 'Neural-network reward in MaxEnt IRL', 'Less interpretable; grid worlds'],
        ['Finn, Levine, Abbeel', 'ICML 2016, New York', 'Guided cost learning, sampled partition function', 'Sampling noise; continuous control'],
        ['McFadden; Prato', 'Econometrics 1974; J. Choice Modelling 2009', 'Logit over a route choice set', 'Needs choice-set generation'],
        ['Dijkstra; Hart et al.', 'Numer. Math. 1959; IEEE SSC 1968', 'Shortest path, fixed metric', 'No learning: our baseline'],
        ['Newson & Krumm', 'ACM SIGSPATIAL 2009, Seattle', 'HMM map matching', 'Our preprocessing step']]
table(s, 0.6, 1.6, 12.1, rows, [2.1, 3.0, 3.8, 3.2], 13, 0.55, hl_row=1)

# ================= 4. HOW WE DIFFER =================
s = new('How this project differs from prior work', 'Literature review',
        "Stress honesty: this is not a new algorithm. The contribution is an open, reproducible, statistically careful study on public data, with an identifiable feature set including goal direction and familiarity, plus an explainable demo. Say where others are stronger: Ziebart has a larger network, Wulfmeier more expressive rewards, Finn continuous control.")
card(s, 0.6, 1.6, 6.0, 4.95)
tb(s, 0.85, 1.72, 5.5, 0.4, 'What we add', 20, True, NAVY, HF)
bullets(s, 0.85, 2.25, 5.5, 4.2, ['Public, reproducible data:: Porto ECML/PKDD 2015 + OpenStreetMap + HMM map matching',
        'Edge-state MDP:: incoming edge as state keeps turn features Markov',
        'Identifiable features:: reference levels, per-step cost, goal direction, edge familiarity',
        'Exact gradients:: soft value iteration, no sampling noise',
        'Honest evaluation:: bootstrap CIs, paired tests, 5-fold CV, ablations',
        'Explainable demo:: per-feature "why this route"'], 15, gap=5)
card(s, 6.9, 1.6, 5.85, 4.95)
tb(s, 7.15, 1.72, 5.4, 0.4, 'Where others are stronger', 20, True, RED, HF)
bullets(s, 7.15, 2.25, 5.4, 4.2, ['Ziebart (AAAI 2008):: original method, much larger road network',
        'Ziebart (UbiComp 2008):: infers an unknown destination, uses context',
        'Wulfmeier (2015):: non-linear rewards',
        'Finn et al. (ICML 2016):: scales to continuous control',
        'Logit models:: standard errors on parameters',
        'We do not claim:: a new algorithm'], 15, gap=5)

# ================= 5. MDP FORMULATION =================
s = new('Our MDP: the course grid-world, translated to roads', 'MDP formulation',
        "Use the course notes as the anchor: tuple S, A, P, R, gamma; terminal states plus and minus 2 in the grid. Ours: state is the incoming edge which fixes node and heading; actions are outgoing roads; transitions are deterministic; reward is w transpose phi, unknown and learned; gamma one with horizon 80. Positive terminal = destination, negative terminal = dead ends (value minus infinity, never chosen).")
# mini grid
gx, gy, cw, ch_ = 0.7, 2.0, 0.72, 0.62
tb(s, 0.6, 1.5, 3.4, 0.35, 'Course example (3×4 grid)', 14, True, NAVY)
for r in range(4):
    for c in range(3):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(gx + c * cw), Inches(gy + r * ch_), Inches(cw), Inches(ch_))
        sh.fill.solid(); sh.line.color.rgb = MUTED; sh.line.width = Pt(1); sh.shadow.inherit = False
        txt, col, f = '−0.5', MUTED, WHITE
        if (r, c) == (0, 2): txt, col, f = '−2', RED, RGBColor(0xFE, 0xE2, 0xE2)
        elif (r, c) == (3, 2): txt, col, f = '+2', GREEN, RGBColor(0xDC, 0xFC, 0xE7)
        elif (r, c) == (1, 1): txt, col, f = '', MUTED, RGBColor(0x94, 0xA3, 0xB8)
        sh.fill.fore_color.rgb = f
        tf = sh.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        rr = tf.paragraphs[0].add_run(); rr.text = txt; rr.font.size = Pt(13); rr.font.bold = True; rr.font.color.rgb = col; rr.font.name = BF
tb(s, 0.6, 4.6, 3.6, 1.3, ['0.6 forward, 0.2 slip each side', 'step reward −0.5, terminals ±2', 'solved by value iteration (max)'], 12, False, MUTED)
rows = [['', 'Course grid-world', 'Our road MDP'],
        ['S', 'Grid cell', 'Incoming directed edge (6,784) + fixed destination d'],
        ['A', 'Up/down/left/right', 'Outgoing roads at the junction (2–4)'],
        ['P', '0.6 / 0.2 / 0.2 slip', 'Deterministic, P = 1 (randomness is in the policy)'],
        ['R', '−0.5 per step, ±2', 'R(e,e′) = wᵀφ(e,e′), w unknown: learned by IRL'],
        ['Terminals', '+2 goal, −2 pit', '+ destination (V = 0)   − dead end (V = −∞, never chosen)'],
        ['γ', '1', '1, finite horizon H = 80'],
        ['V, π', 'V = R + max Σ P V′', 'V(e) = log Σ exp[R + V(e′)],  π = exp(Q − V)']]
table(s, 4.5, 1.55, 8.3, rows, [1.2, 2.3, 4.8], 13, 0.62)

# ================= 6. MDP DIAGRAM =================
s = new('Worked example: states, actions, terminals', 'MDP formulation',
        "Read the diagram left to right. Circles are states (incoming road edges), diamonds are actions (choose an outgoing road), arrows from diamonds carry P=1 and the reward. Green double circle is the destination: the positive terminal. Red is a dead end: the negative terminal with value minus infinity so the policy never enters it. At S1 the model prefers the straight arterial (0.68), sometimes the side street (0.31), almost never the U-turn (0.01). The +5 goal bonus is a constant per trip, so the policy is unchanged. Toy topology; the weights are the real learned ones.")
s.shapes.add_picture('mdp_diagram.png', Inches(1.85), Inches(1.4), height=Inches(5.85))

# ================= 7. REWARD / FEATURES =================
s = new('The reward: 11 interpretable features', 'Method',
        "These are the learned weights from the final v3 model (single split). Sharp turns are strongly avoided, familiar edges strongly preferred, hop cost and slight turns penalised, moving away from the goal penalised. Note length and travel time are correlated (0.81), so read those two together; the earlier feature set was unidentifiable because of collinear one-hots, which we fixed with reference levels and a per-step constant.")
names = ['edge_familiarity', 'motorway', 'travel_time', 'num_lanes', 'same_road_name', 'primary', 'length', 'away_from_goal', 'turn_slight', 'step', 'turn_sharp']
vals = [0.872, 0.230, 0.157, -0.008, -0.042, -0.111, -0.235, -0.295, -0.519, -0.674, -1.464]
cols = [GREEN if v > 0 else RED for v in vals]
bar_chart(s, 0.6, 1.5, 7.3, 5.4, names, vals, 'Learned weights w (v3, scaled features)', True, cols, '0.00', -1.8, 1.2, 12)
card(s, 8.2, 1.6, 4.55, 5.2)
tb(s, 8.4, 1.7, 4.2, 0.35, 'What it tells us', 18, True, NAVY, HF)
bullets(s, 8.4, 2.2, 4.2, 4.5, ['Sharp turns / U-turns:: strongest penalty (−1.46)', 'Familiar edges:: strongest bonus (+0.87): drivers use roads other taxis use',
        'Per-step cost:: −0.67 prefers fewer segments', 'Slight turns:: −0.52, straight is the reference', 'Away from goal:: −0.29, goal-directed',
        'Identifiable:: reference levels + step constant fix v1 collinearity'], 14, gap=5)

# ================= 8. ALGORITHM =================
s = new('Maximum Entropy IRL with exact soft value iteration', 'Method',
        "Outer loop: L-BFGS on weights w. Inner loop per destination: backward pass is soft value iteration, giving V and Q — dynamic programming control. Forward pass computes expected state visitation — dynamic programming prediction. Gradient is expected minus expert feature counts plus L2. Exact because the MDP is small, discrete and deterministic, no sampling. Then the most likely route is extracted with Viterbi/Dijkstra on V minus Q. Training about 13 minutes per CV fold on CPU.")
steps = [('1', 'Backward pass', 'Soft value iteration\nV(e) = log Σ exp[R + V(e′)]\nQ, π = exp(Q − V)', BLUE),
         ('2', 'Forward pass', 'Expected state-visit\nfrequencies under π\n→ E_π[φ]', TEAL),
         ('3', 'Gradient', '∇ = E_π[φ] − φ_expert + λw\n(exact, no sampling)', AMBER),
         ('4', 'L-BFGS-B', 'Update w, repeat\n(~25–35 iterations)', NAVY)]
for i, (n_, h, d, c) in enumerate(steps):
    x = 0.6 + i * 3.1
    card(s, x, 1.7, 2.85, 2.55)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(1.85), Inches(0.5), Inches(0.5)); o.fill.solid(); o.fill.fore_color.rgb = c; o.line.fill.background(); o.shadow.inherit = False
    tf = o.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER; r = tf.paragraphs[0].add_run(); r.text = n_; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = BF
    tb(s, x + 0.75, 1.88, 2.0, 0.45, h, 16, True, NAVY, HF)
    tb(s, x + 0.15, 2.55, 2.6, 1.6, d.split('\n'), 13, False, TXT)
card(s, 0.6, 4.55, 12.15, 2.2)
bullets(s, 0.85, 4.7, 5.7, 2.0, ['Trajectory probability:: P(τ|w) = exp(wᵀφ(τ)) / Z', 'Route prediction:: Viterbi/Dijkstra on V(e) − Q(e,e′)', 'Calibration:: extra length penalty α chosen on a validation slice, never on test'], 14, gap=5)
bullets(s, 6.8, 4.7, 5.8, 2.0, ['Why exact:: small discrete MDP, so DP replaces sampling', 'Verified:: finite-difference gradient test, relative error ~1e-8', 'Cost:: ~13 min per fold on 8 CPU cores'], 14, gap=5)

# ================= 9. DATA & PIPELINE =================
s = new('Data pipeline: from taxi GPS to road edges', 'Method',
        "Porto ECML/PKDD 2015 taxi dataset, 1.7M raw trips; we filtered to central Porto and map-matched with an HMM (Leuven MapMatching): only 40% of attempted trips matched, 3,000 used. Graph from OSM, 3,367 nodes, 6,784 directed edges. Be upfront: matching errors are treated as ground truth. Familiarity counts come from a separate disjoint pool of trips to avoid leakage.")
flow = [('Porto taxi GPS', 'ECML/PKDD 2015\ncentral bbox filter'), ('HMM map matching', '7,528 tried\n3,000 matched (40%)'), ('Edge sequences', '5–75 edges per trip\nwith context'), ('MaxEnt IRL', '11 features\nexact gradients'), ('Routes + demo', 'evaluation, API,\nexplainable dashboard')]
for i, (h, d) in enumerate(flow):
    x = 0.6 + i * 2.5
    card(s, x, 1.7, 2.3, 1.8)
    tb(s, x + 0.1, 1.82, 2.1, 0.5, h, 15, True, NAVY, HF, PP_ALIGN.CENTER)
    tb(s, x + 0.1, 2.45, 2.1, 0.95, d.split('\n'), 12.5, False, TXT, BF, PP_ALIGN.CENTER)
    if i < 4:
        tb(s, x + 2.28, 2.3, 0.3, 0.5, '→', 22, True, TEAL, BF, PP_ALIGN.CENTER)
stat(s, 0.6, 3.9, 2.9, '6,784', 'directed edges', '3,367 nodes, central Porto', NAVY)
stat(s, 3.7, 3.9, 2.9, '3,000', 'matched trips', 'of 7,528 attempted', TEAL)
stat(s, 6.8, 3.9, 2.9, '2,040 / 360 / 600', 'fit / val / test per fold', '5-fold cross-validation', BLUE, 24)
stat(s, 9.9, 3.9, 2.85, '600', 'held-out test trips', 'single-split CIs use these', AMBER)
tb(s, 0.6, 5.7, 12.1, 1.0, ['Edge familiarity = log(1 + #trips through the edge) from an independent pool of map-matched trips, disjoint from train, validation and test.',
   'Map-matching errors are treated as ground truth (a limitation).'], 14, False, MUTED)

# ================= 10. SETUP & METRICS =================
s = new('Baselines and metrics', 'Experimental setup',
        "Four baselines. Dijkstra shortest is the strong one on taxi data. Pop-Shortest is a popularity-discounted shortest path — a simple non-IRL baseline that turned out to tie us in cross-validation, which we report honestly. Metrics: edge Jaccard and F1 for route overlap; top-k junction accuracy and log-likelihood for the probabilistic model; bootstrap CIs and paired Wilcoxon tests.")
card(s, 0.6, 1.6, 5.9, 5.0)
tb(s, 0.85, 1.7, 5.4, 0.4, 'Baselines', 20, True, NAVY, HF)
bullets(s, 0.85, 2.25, 5.4, 4.3, ['Dijkstra-Shortest:: minimise length', 'Dijkstra-Fastest:: minimise free-flow time', 'Myopic classifier:: per-junction logistic regression, no destination',
        'Pop-Shortest:: shortest path discounted by edge popularity'], 18, gap=12)
card(s, 6.8, 1.6, 5.95, 5.0)
tb(s, 7.05, 1.7, 5.4, 0.4, 'Metrics', 20, True, NAVY, HF)
bullets(s, 7.05, 2.25, 5.5, 4.3, ['Edge F1 / Jaccard:: overlap of predicted and driven edges', 'Path length ratio:: detects loops or shortcuts', 'Top-1 / Top-3 accuracy:: next-road prediction at each junction',
        'Log-likelihood / trip NLL:: how well the model explains held-out trips', 'Evidence:: 95% bootstrap CIs, paired Wilcoxon tests, 5-fold CV'], 17, gap=12)

# ================= 11. RESULTS HEADLINE =================
s = new('Results: learned routes beat the shortest path', 'Results',
        "Single split, 600 held-out Porto trips. Calibrated IRL v3 gets edge F1 0.682 (95% CI 0.657-0.708) against 0.636 for Dijkstra shortest; paired p = 2.7e-5. Raw v3 without calibration is 0.659 (p = 0.06). Fastest is 0.454, myopic 0.162. Also point out the v1 original features lost to shortest (0.611) — that is why we redesigned the features.")
names = ['IRL-v3 (calibrated)', 'IRL-v3 (raw)', 'IRL-v2 (calibrated)', 'Dijkstra-Shortest', 'IRL-v1 (original feats)', 'Dijkstra-Fastest', 'Myopic classifier']
vals = [0.682, 0.659, 0.648, 0.636, 0.611, 0.454, 0.162]
cols = [TEAL, TEAL, TEAL, AMBER, TEAL, MUTED, MUTED]
bar_chart(s, 0.6, 1.5, 7.6, 5.4, names, vals, 'Edge F1 on 600 held-out trips', True, cols, '0.000', 0, 0.8, 12)
stat(s, 8.5, 1.6, 4.25, '0.682', 'Edge F1 (IRL-v3)', '95% CI 0.657–0.708', TEAL)
stat(s, 8.5, 3.3, 2.05, '+0.046', 'vs Shortest', 'paired p = 2.7e-5', GREEN)
stat(s, 10.7, 3.3, 2.05, '93.7%', 'Top-1 acc.', 'Top-3: 99.9%', BLUE)
tb(s, 8.5, 5.0, 4.25, 1.8, ['Trip NLL 5.46 (v1: 6.12)', 'Log-likelihood −0.155 per step', 'Path length ratio 0.91 vs 0.88 for shortest'], 14, False, TXT)

# ================= 12. CROSS-VALIDATION =================
s = new('Cross-validation: the gain holds, but a baseline ties', 'Results',
        "5-fold CV, mean over folds. IRL v3 calibrated 0.669 plus or minus 0.010 versus Dijkstra shortest 0.630 plus or minus 0.003. But Pop-Shortest, a simple popularity-weighted shortest path, scores 0.671 — it ties us. So the honest reading: much of the improvement comes from the familiarity signal, and IRL's added value is a calibrated probabilistic route model with interpretable weights, not a higher F1.")
names = ['Dijkstra-Fastest', 'Dijkstra-Shortest', 'IRL-v3 (raw)', 'IRL-v3 (calibrated)', 'Pop-Shortest']
vals = [0.478, 0.630, 0.655, 0.669, 0.671]
bar_chart(s, 0.6, 1.5, 6.9, 4.6, names, vals, 'Edge F1, 5-fold CV (mean)', True, [MUTED, AMBER, TEAL, TEAL, BLUE], '0.000', 0, 0.8, 12)
rows = [['Method', 'Edge F1', 'Jaccard'], ['Dijkstra-Fastest', '0.478 ± 0.009', '0.379 ± 0.008'], ['Dijkstra-Shortest', '0.630 ± 0.003', '0.538 ± 0.003'],
        ['IRL-v3 (raw)', '0.655 ± 0.008', '0.572 ± 0.010'], ['IRL-v3 (calibrated)', '0.669 ± 0.010', '0.588 ± 0.011'], ['Pop-Shortest', '0.671 ± 0.007', '0.594 ± 0.005']]
table(s, 7.8, 1.6, 4.95, rows, [2.0, 1.5, 1.45], 12.5, 0.5, hl_row=4)
card(s, 7.8, 4.85, 4.95, 1.75, fill=RGBColor(0xFE, 0xF3, 0xC7), line=RGBColor(0xF5, 0x9E, 0x0B))
tb(s, 7.95, 4.95, 4.7, 0.35, 'Honest takeaway', 14, True, NAVY, HF)
tb(s, 7.95, 5.35, 4.7, 1.2, 'IRL beats shortest path by ~0.04 F1, but Pop-Shortest matches it. Our value: a calibrated route distribution with readable weights.', 14, False, TXT)
tb(s, 0.6, 6.2, 6.9, 0.6, 'Mean ± std over 5 folds (one seed). Calibration = extra length penalty α chosen on validation data.', 11, False, MUTED)

# ================= 13. FEATURE PROGRESSION =================
s = new('What made the model work: feature-set progression', 'Results',
        "This is a real ablation of the modelling choices. v1 original features lost to Dijkstra. v2 made the parameterisation identifiable and added goal-direction; the likelihood improved only slightly. Calibration with a length penalty lifted F1. v3 added edge familiarity, which gave the largest likelihood gain — NLL 6.06 to 5.46 — and the best F1.")
cd = CategoryChartData(); cd.categories = ['v1', 'v2 raw', 'v2 calibrated', 'v3 raw', 'v3 calibrated']; cd.add_series('Edge F1', [0.611, 0.606, 0.648, 0.659, 0.682])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5), Inches(7.2), Inches(5.0), cd); ch = gf.chart
ch.has_legend = False; ch.has_title = True; ch.chart_title.text_frame.text = 'Edge F1 by feature set (600 test trips)'
r = ch.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.bold = True; r.font.name = BF; r.font.color.rgb = NAVY
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.number_format = '0.000'; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(12)
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.minimum_scale = 0.5; ch.value_axis.maximum_scale = 0.72; ch.value_axis.major_gridlines.format.line.color.rgb = LINE
ch.value_axis.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.size = Pt(12)
ch.value_axis.tick_labels.number_format = '0.00'; ch.value_axis.tick_labels.number_format_is_linked = False
for i, c in enumerate([MUTED, MUTED, TEAL, TEAL, GREEN]):
    p = pl.series[0].points[i]; p.format.fill.solid(); p.format.fill.fore_color.rgb = c
rows = [['Model', 'Trip NLL ↓', 'Top-1'], ['v1', '6.12', '92.7%'], ['v2', '6.06', '92.9%'], ['v3', '5.46', '93.7%']]
table(s, 8.1, 1.65, 4.65, rows, [1.6, 1.6, 1.45], 13, 0.5, hl_row=3)
bullets(s, 8.1, 3.9, 4.65, 2.8, ['v1:: original features lost to Dijkstra (0.611 vs 0.636, p = 0.029)', 'v2:: identifiable parameterisation + goal direction', 'v3:: + edge familiarity: biggest likelihood gain'], 14, gap=6)

# ================= 14. ABLATIONS =================
s = new('Ablations: familiarity and turns matter most', 'Results',
        "Five feature-group ablations, each retrained from scratch on the same 600 test trips. Removing familiarity costs the most F1 (-0.034) and NLL (+0.60). Removing turn features collapses the likelihood (NLL +1.78) yet F1 slightly rises: validation chose a bigger length penalty for that model; route overlap and likelihood disagree there. Goal direction and road class add almost nothing. Caveat: single split, one seed; F1 CI is about plus or minus 0.025, so only familiarity clearly exceeds noise; NLL is the cleaner signal.")
cats = ['no familiarity', 'no step cost', 'no goal direction', 'no road class', 'no turn features']
bar_chart(s, 0.6, 1.5, 6.0, 4.3, cats, [-0.034, -0.016, -0.001, 0.001, 0.011], 'Change in edge F1 vs full model (0.682)', True, [RED, RED, MUTED, MUTED, AMBER], '+0.000;-0.000', -0.05, 0.03, 12)
bar_chart(s, 6.8, 1.5, 6.0, 4.3, cats, [0.60, 0.25, 0.04, 0.05, 1.78], 'Change in trip NLL (higher = worse), full = 5.46', True, [RED, RED, MUTED, MUTED, RED], '+0.00', 0, 2.2, 12)
card(s, 0.6, 5.95, 12.15, 0.85, fill=RGBColor(0xFE, 0xF3, 0xC7), line=AMBER)
tb(s, 0.8, 6.0, 11.8, 0.75, 'Turn features: likelihood collapses (NLL +1.78) but F1 rises, because validation picked a larger length penalty (α = 1.0). Overlap and likelihood disagree here. Single split, one seed: the F1 95% CI is about ±0.025.', 13, False, TXT, anchor=MSO_ANCHOR.MIDDLE)

# ================= 15. DEMO =================
s = new('Live demo: compare routes and ask "why this route?"', 'Demo',
        "Demo script: open localhost:5000. Click Random trip or type a street name like Avenida da Boavista as start and Rua da Ribeira Negra as destination. Four routes appear: learned, shortest, fastest, and the real driver (dashed green). Click the learned route to open the Why panel: the reward is linear, so each feature's contribution is exact — e.g. fewer segments, motorway use and slight turns. Caveat: this explains the model, not the driver's intent. Names come from the OSM graph; landmark search would need an external geocoder.")
s.shapes.add_picture('demo_routes.jpg', Inches(0.6), Inches(1.6), height=Inches(2.95))
s.shapes.add_picture('demo_why.jpg', Inches(5.45), Inches(1.6), height=Inches(2.95))
card(s, 0.6, 4.95, 12.15, 1.85)
bullets(s, 0.85, 5.05, 5.8, 1.7, ['Compare:: learned vs shortest vs fastest vs the real driver', 'Search:: type street names (offline, from the OSM graph) or click the map'], 14, gap=5)
bullets(s, 6.8, 5.05, 5.8, 1.7, ['Explain:: exact per-feature contributions, since R = Σ weight × feature', 'Replay:: held-out real taxi trips, with overlap % to the driver'], 14, gap=5)
card(s, 9.85, 1.55, 2.9, 3.15, fill=RGBColor(0xEC, 0xF3, 0xFB), line=RGBColor(0xBF, 0xD4, 0xEE))
tb(s, 10.0, 1.65, 2.6, 0.4, 'Run it', 16, True, NAVY, HF)
tb(s, 10.0, 2.1, 2.65, 2.5, ['python run_api.py', '', 'open http://localhost:5000', '', 'Uses the v3 model and', 'validation-tuned α'], 12.5, False, TXT, 'Consolas')

# ================= 16. LIMITATIONS =================
s = new('Limitations and honest takeaways', 'Discussion',
        "Be upfront: modest gain over strong baselines, popularity baseline ties us, single seed, one city, one reward for all drivers, map matching noise. The deck's value is a reproducible, interpretable benchmark and a clear picture of what drives route choice.")
items = [('Modest margin', 'IRL beats shortest by ~0.04 F1; a popularity-weighted shortest path ties it in CV.', RED),
         ('Single seed, one city', 'Fold variation reported; seed variation and other cities not measured.', RED),
         ('Linear, single reward', 'No feature interactions, one w for every driver and time of day.', AMBER),
         ('Map-matching noise', 'Only 40% of trips matched; errors are treated as ground truth.', AMBER),
         ('Static network', 'No congestion or time-of-day in the state; free-flow times only.', AMBER),
         ('Model-level explanations', 'They show what the reward values, not why a driver chose a road.', AMBER)]
for i, (h, d, c) in enumerate(items):
    x = 0.6 + (i % 3) * 4.1; y = 1.65 + (i // 3) * 2.55
    card(s, x, y, 3.9, 2.3)
    tb(s, x + 0.2, y + 0.15, 3.5, 0.5, h, 17, True, c, HF)
    tb(s, x + 0.2, y + 0.75, 3.5, 1.4, d, 14, False, TXT)

# ================= 17. CONCLUSION =================
s = new('Conclusion and future work', 'Wrap-up',
        "Summarise in three sentences: we learned an interpretable reward from real taxi data; it beats shortest path significantly but modestly and is tied by a popularity baseline; the main value is the exact, reproducible, explainable pipeline. Future: seeds and sample curves, nonlinear/context rewards, more cities, unknown destination.")
card(s, 0.6, 1.6, 6.0, 5.1)
tb(s, 0.85, 1.7, 5.5, 0.4, 'What we showed', 20, True, NAVY, HF)
bullets(s, 0.85, 2.25, 5.5, 4.3, ['Real data:: 3,000 Porto taxi trips, HMM map-matched', 'Edge-state MDP:: exact MaxEnt IRL, 11 interpretable features', 'F1 0.682 vs 0.636:: significant (p = 2.7e-5), CV 0.669 vs 0.630',
        'Drivers prefer:: familiar roads, few turns, goal-directed moves', 'Open demo:: explainable route comparison'], 17, gap=11)
card(s, 6.9, 1.6, 5.85, 5.1)
tb(s, 7.15, 1.7, 5.4, 0.4, 'Future work', 20, True, NAVY, HF)
bullets(s, 7.15, 2.25, 5.4, 4.3, ['More seeds and sample-size curves::', 'Context-aware reward:: time of day, driver clusters', 'Non-linear reward:: neural features, keep explanations', 'Other cities and larger graphs::', 'Unknown destination:: infer it from a partial trip', 'Forward RL:: use the learned reward for navigation or AV policy training'], 17, gap=11)

# ================= 18. REFERENCES =================
s = new('References', 'Appendix', "Verify venue and page details against the originals before final submission.")
refs = ['[1] B. D. Ziebart, A. Maas, J. A. Bagnell, A. K. Dey, "Maximum entropy inverse reinforcement learning," AAAI, 2008, pp. 1433–1438.',
        '[2] B. D. Ziebart, A. Maas, A. K. Dey, J. A. Bagnell, "Navigate like a cabbie: probabilistic reasoning from observed context-aware behavior," UbiComp, 2008.',
        '[3] A. Y. Ng, S. J. Russell, "Algorithms for inverse reinforcement learning," ICML, 2000, pp. 663–670.',
        '[4] P. Abbeel, A. Y. Ng, "Apprenticeship learning via inverse reinforcement learning," ICML, 2004.',
        '[5] M. Wulfmeier, P. Ondruska, I. Posner, "Maximum entropy deep inverse reinforcement learning," arXiv:1507.04888, 2015.',
        '[6] C. Finn, S. Levine, P. Abbeel, "Guided cost learning: deep inverse optimal control via policy optimization," ICML, 2016.',
        '[7] D. McFadden, "Conditional logit analysis of qualitative choice behavior," Frontiers in Econometrics, 1974.',
        '[8] P. Newson, J. Krumm, "Hidden Markov map matching through noise and sparseness," ACM SIGSPATIAL GIS, 2009.',
        '[9] L. Moreira-Matias et al., "Predicting taxi–passenger demand using streaming data," IEEE Trans. ITS, 14(3), 2013.',
        '[10] G. Boeing, "OSMnx: new methods for acquiring, constructing, analyzing, and visualizing complex street networks," Comput. Environ. Urban Syst., 65, 2017.']
tb(s, 0.6, 1.5, 12.1, 5.5, refs, 15, False, TXT)

prs.save('Route_Preference_IRL_Final.pptx')
print('saved', N[0], 'slides')
